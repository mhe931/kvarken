"""Export this canonical LaTeX manuscript using the official Word template."""

import argparse
import re
import subprocess
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

WORKFLOW_FIGURE = "Intended student workflow for using the EO infrastructure with CSC"
FIGURES = (
    "Provider-to-application architecture",
    "Sentinel-1A VV gamma-zero input",
    "Sentinel-2A Level-2A false-colour observation",
    "Selective Sentinel-2 optical proxy labels",
    "Scene-specific SAR candidate classes derived",
    "Held-out agreement between thesis-derived",
)
FIGURE_DESCRIPTIONS = (
    "Vector architecture diagram showing the metadata path from provider through transport "
    "and validation, typed EOScene, provenance and SQLite catalog, read-only queries/API/"
    "reports, and the scientific application. The case-study raster path used local "
    "Planetary Computer COG assets; CSC execution covered metadata only.",
    "Sentinel-1A VV gamma-zero radar backscatter in eastern Kvarken on 30 March 2021. "
    "Grey tones show measured backscatter, not an ice classification; patterns can also "
    "reflect wind, incidence angle and residual preprocessing effects.",
    "Independent Sentinel-2A false-colour optical observation of eastern Kvarken, "
    "approximately 17.90 hours after the radar acquisition. Bright snow-covered ice "
    "and dark water support selective proxy labels, not field ground truth. "
    "Grey areas lack usable coverage; dark pixels are not a missing-data mask.",
    "Selective optical proxy reference on the analysed water domain: pale-cyan snow-ice "
    "labels, blue dark-water labels, grey water without a label and white excluded "
    "land/coastal areas. "
    "The proxy does not label all ice types or establish total sea-ice extent.",
    "Scene-specific SAR candidate map, with pale-cyan ice candidates and blue water "
    "candidates within the screened domain; white areas are excluded. "
    "Candidate ice covers 1,219.28 square "
    "kilometres; 65.28 percent has no optical label. This is not validated total ice extent.",
    "Held-out spatial-block agreement map comparing SAR candidates with selective "
    "optical labels. Green cells agree and red cells disagree; other cells are "
    "unlabelled or excluded. Ice precision is 23.53 percent against this proxy.",
)
TABLES = (
    ("Comparison of EO platforms, standards, and software ecosystems", "2.1"),
    ("Research questions, primary evidence,", "3.1"),
    ("Frozen 256-scene synthetic experiment baseline.", "6.1"),
    ("Held-out confusion matrix.", "6.2"),
    ("Held-out proxy agreement.", "6.3"),
    ("Evidence status for the proposed student workflow", "9.1"),
)
WORKFLOW_STAGES = (
    "Student onboarding",
    "CSC environment and access setup",
    "EO Catalogue discovery",
    "Dataset selection",
    "Data retrieval or remote access",
    "Processing and analysis",
    "Provenance recording",
    "Result storage",
    "Reproduction by another student",
)


def _field(paragraph, instruction: str) -> None:
    for tag in ("begin", "separate", "end"):
        run = paragraph.add_run()
        marker = OxmlElement("w:fldChar")
        marker.set(qn("w:fldCharType"), tag)
        run._r.append(marker)
        if tag == "begin":
            text = OxmlElement("w:instrText")
            text.set(qn("xml:space"), "preserve")
            text.text = " " + instruction + " "
            run._r.append(text)
        elif tag == "separate":
            paragraph.add_run("Update field")


def _prefix(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, run._r)


def _remove_repeated_workflow_table_header(document: Document) -> None:
    table = None
    for caption in document.paragraphs:
        if not caption.text.startswith(
            "Table 9.1. Evidence status for the proposed student workflow"
        ):
            continue
        element = caption._p.getnext()
        while (
            element is not None
            and element.tag == qn("w:p")
            and not Paragraph(element, document).text.strip()
        ):
            element = element.getnext()
        if element is not None and element.tag == qn("w:tbl"):
            table = next((table for table in document.tables if table._tbl is element), None)
            if table is not None:
                break
    if table is None or len(table.rows) < 2:
        return
    expected_header = ("Workflow stage", "Status", "Evidence and boundary")
    first_row = tuple(cell.text.strip() for cell in table.rows[0].cells)
    second_row = tuple(cell.text.strip() for cell in table.rows[1].cells)
    if first_row == second_row == expected_header:
        table._tbl.remove(table.rows[1]._tr)


def _correct_appendix_cross_references(document: Document) -> None:
    for element in document.element.xpath(".//w:p"):
        paragraph = Paragraph(element, document)
        if not re.search(r"Appendix\s+12\b", paragraph.text):
            continue
        preceding_text = ""
        for child in paragraph._p:
            if child.tag == qn("w:hyperlink"):
                text_nodes = list(child.iter(qn("w:t")))
                if "".join(node.text or "" for node in text_nodes) == "12" and re.search(
                    r"Appendix\s*$", preceding_text
                ):
                    text_nodes[-1].text = "B"
                    break
            preceding_text += "".join(node.text or "" for node in child.iter(qn("w:t")))


def format_document(path: Path) -> dict:
    document = Document(path)
    document.core_properties.title = (
        "Design and Implementation of a Hybrid Earth Observation Data Infrastructure"
    )
    document.core_properties.author = "Däniel Ebrahimzadeh Esfahani"
    document.core_properties.last_modified_by = "AI-assisted manuscript export"
    paragraphs = document.paragraphs
    abstract_start = next(
        (i for i, p in enumerate(paragraphs) if p.text.startswith("UNIVERSITY OF VAASA")), None
    )
    abstract_end = next(
        (i for i, p in enumerate(paragraphs) if p.text.startswith("Keywords:")), None
    )
    if abstract_start is None or abstract_end is None or abstract_start >= abstract_end:
        raise ValueError("canonical abstract and keywords are required")
    for i, paragraph in enumerate(paragraphs):
        drawings = paragraph._p.xpath(".//wp:docPr")
        if not drawings:
            continue
        caption = next((p.text for p in paragraphs[i + 1 :] if p.text.strip()), "")
        matched = False
        for prefix, description in zip(FIGURES[1:], FIGURE_DESCRIPTIONS[1:], strict=True):
            if caption.startswith(prefix):
                for drawing in drawings:
                    drawing.set("descr", description)
                matched = True
                break
        if not matched and i < abstract_start:
            for drawing in drawings:
                drawing.set("descr", "University of Vaasa logo")
    styles = document.styles
    for name in ("Normal", "Body Text", "Bibliography"):
        if name not in styles:
            continue
        styles[name].font.name = "Calibri"
        styles[name].font.size = Pt(12)
        styles[name].paragraph_format.line_spacing = 1.5
        styles[name].paragraph_format.space_after = Pt(6)
    for level, size in ((1, 16), (2, 14), (3, 12)):
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.page_break_before = level == 1
        style.paragraph_format.keep_with_next = True
        if style.element.pPr is not None:
            number = style.element.pPr.find(qn("w:numPr"))
            if number is not None:
                style.element.pPr.remove(number)
    front = styles.add_style("Front Matter", 1)
    front.base_style = styles["Heading 1"]
    for section in document.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(3)
        section.left_margin = section.right_margin = Cm(3)
        section.header_distance = Cm(1.5)
        section.different_first_page_header_footer = True
        page_numbering = section._sectPr.find(qn("w:pgNumType"))
        if page_numbering is None:
            page_numbering = OxmlElement("w:pgNumType")
            columns = section._sectPr.find(qn("w:cols"))
            if columns is None:
                section._sectPr.append(page_numbering)
            else:
                section._sectPr.insert(section._sectPr.index(columns), page_numbering)
        page_numbering.set(qn("w:start"), "1")
        header = section.header.paragraphs[0]
        header.clear()
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(header, "PAGE")
        for run in header.runs:
            run.font.name, run.font.size = "Calibri", Pt(11)
        for footer in (section.footer, section.first_page_footer, section.first_page_header):
            for paragraph in footer.paragraphs:
                paragraph.clear()
    for i, paragraph in enumerate(paragraphs):
        if i < abstract_start:
            if not paragraph.text and not paragraph._p.xpath(".//w:drawing"):
                paragraph._p.getparent().remove(paragraph._p)
                continue
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1
            suppress = OxmlElement("w:suppressAutoHyphens")
            paragraph._p.get_or_add_pPr().append(suppress)
            if paragraph.text.startswith(("Master", "School", "Major", "Supervisor")):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            size = 12
            if paragraph.text == "Däniel Ebrahimzadeh Esfahani":
                size = 16
                paragraph.paragraph_format.space_before = Pt(81)
            elif paragraph.text.startswith("Design and Implementation"):
                size = 20
                paragraph.paragraph_format.space_before = Pt(20)
            elif paragraph.text.startswith("for Research"):
                size = 14
                paragraph.paragraph_format.space_after = Pt(35)
            elif paragraph.text == "Vaasa 2026":
                paragraph.paragraph_format.space_before = Pt(230)
            for run in paragraph.runs:
                run.font.name, run.font.size = "Calibri", Pt(size)
                run.bold = paragraph.text.startswith("Design and Implementation")
        elif i <= abstract_end:
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if paragraph.text.startswith(("UNIVERSITY", "Keywords:", "Abstract"))
                else WD_ALIGN_PARAGRAPH.JUSTIFY
            )
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name, run.font.size = "Calibri", Pt(11)
    paragraphs[abstract_start].paragraph_format.page_break_before = True
    if document.tables:
        for row in document.tables[0].rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.name, run.font.size = "Calibri", Pt(11)
    for table in document.tables:
        if not table.rows:
            continue
        if table.cell(0, 0).text != "ACID":
            continue
        for header in table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"):
            header.getparent().remove(header)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.space_after = Pt(3)
    chapter = section = subsection = figure_number = 0
    appendix = False
    appendix_letter = ""
    figure_entries, table_entries = [], []
    bookmark_id = (
        max(
            (int(node.get(qn("w:id"))) for node in document.element.xpath(".//w:bookmarkStart")),
            default=0,
        )
        + 1
    )
    for paragraph in paragraphs[abstract_end + 1 :]:
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.paragraph_format.keep_with_next = True
        text = paragraph.text
        if paragraph.style.name == "Heading 1":
            if text in {"Acknowledgements", "Abbreviations"}:
                paragraph.style = front
            elif text == "Disclosure of AI-Assisted Work":
                appendix = True
                appendix_letter = "A"
                section = subsection = 0
                paragraph.style = front
                _prefix(paragraph, "A  ")
            elif appendix:
                appendix_letter = "B"
                section = subsection = 0
                paragraph.style = front
                _prefix(paragraph, "B  ")
            else:
                chapter += 1
                section = subsection = 0
                _prefix(paragraph, f"{chapter}  ")
        elif paragraph.style.name == "Heading 2":
            section += 1
            subsection = 0
            _prefix(paragraph, f"{appendix_letter or chapter}.{section}  ")
        elif paragraph.style.name == "Heading 3":
            subsection += 1
            _prefix(paragraph, f"{appendix_letter or chapter}.{section}.{subsection}  ")
        architecture_caption = text.startswith(("Figure 4.1. " + FIGURES[0], FIGURES[0]))
        workflow_caption = text.startswith(WORKFLOW_FIGURE)
        figure_caption = architecture_caption or workflow_caption or text.startswith(FIGURES[1:])
        table_number = next((number for prefix, number in TABLES if text.startswith(prefix)), None)
        kind = "Figure" if figure_caption else "Table" if table_number is not None else None
        if kind:
            entries = figure_entries if kind == "Figure" else table_entries
            if kind == "Figure":
                number = (
                    "4.1"
                    if architecture_caption
                    else "9.1"
                    if workflow_caption
                    else f"6.{figure_number + 1}"
                )
            elif table_number is not None:
                number = table_number
            else:
                raise RuntimeError("recognized table caption has no assigned number")
            if kind == "Figure" and not architecture_caption and not workflow_caption:
                figure_number += 1
            if not text.startswith(f"{kind} {number}."):
                _prefix(paragraph, f"{kind} {number}. ")
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.keep_together = True
            if kind == "Table":
                paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.font.size = Pt(11)
            bookmark = f"thesis{kind}{bookmark_id}"
            start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
            start.set(qn("w:id"), str(bookmark_id))
            start.set(qn("w:name"), bookmark)
            end.set(qn("w:id"), str(bookmark_id))
            paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
            paragraph._p.append(end)
            caption_text = text
            numbered_prefix = f"{kind} {number}."
            if caption_text.startswith(numbered_prefix):
                caption_text = caption_text[len(numbered_prefix) :].lstrip()
            short_caption = re.split(r"\.(?:\s|$)", caption_text, maxsplit=1)[0]
            entries.append((f"{kind} {number}. {short_caption}", bookmark))
            bookmark_id += 1
    appendix_start = next(
        p for p in document.paragraphs if p.text == "A  Disclosure of AI-Assisted Work"
    )
    references = appendix_start.insert_paragraph_before("References", "Front Matter")
    for paragraph in list(document.paragraphs):
        if paragraph.style.name == "Bibliography":
            appendix_start._p.addprevious(paragraph._p)
    if references.text != "References":
        raise RuntimeError("references heading was not retained")
    anchor = next(p for p in document.paragraphs if p.text == "Acknowledgements")
    contents = anchor.insert_paragraph_before("Contents", "Front Matter")
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "9")
    contents._p.get_or_add_pPr().append(outline)
    _field(anchor.insert_paragraph_before(), r'TOC \o "1-2" \h \z \u')
    for title, entries in (("Figures", figure_entries), ("Tables", table_entries)):
        heading = anchor.insert_paragraph_before(title, "Front Matter")
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "9")
        heading._p.get_or_add_pPr().append(outline)
        for text, bookmark in entries:
            paragraph = anchor.insert_paragraph_before(text + "\t")
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(15), WD_TAB_ALIGNMENT.RIGHT)
            _field(paragraph, "PAGEREF " + bookmark + r" \h")
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    document.settings.element.append(update)
    _correct_appendix_cross_references(document)
    _remove_repeated_workflow_table_header(document)
    document.save(path)
    return {"figures": len(figure_entries), "tables": len(table_entries)}


def _flow_diagram_table(
    document: Document,
    stages: tuple[str | tuple[str, str], ...],
    caption: str,
    description: str,
):
    table = document.add_table(rows=2 * len(stages) - 1, cols=1)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Cm(15)
    for index, row in enumerate(table.rows):
        cell = row.cells[0]
        cell.width = Cm(15)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        if index == 0 and caption.startswith(f"Figure 9.1. {WORKFLOW_FIGURE}"):
            paragraph.paragraph_format.page_break_before = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.space_before = Pt(2)
        if index % 2:
            run = paragraph.add_run("↓")
            run.font.size = Pt(12)
            continue
        stage = stages[index // 2]
        title, detail = stage if isinstance(stage, tuple) else (stage, "")
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "DDEBF7")
        cell._tc.get_or_add_tcPr().append(shading)
        title_run = paragraph.add_run(title)
        title_run.bold = True
        title_run.font.size = Pt(10)
        if detail:
            title_run.add_break()
            detail_run = paragraph.add_run(detail)
            detail_run.font.size = Pt(10)
    caption_property = OxmlElement("w:tblCaption")
    caption_property.set(qn("w:val"), caption)
    description_property = OxmlElement("w:tblDescription")
    description_property.set(qn("w:val"), description)
    table._tbl.tblPr.append(caption_property)
    table._tbl.tblPr.append(description_property)
    return table


def _insert_tikz_diagrams(path: Path) -> None:
    document = Document(path)
    placeholders = [table for table in document.tables if not table.rows]
    if len(placeholders) != 2:
        raise ValueError(f"expected two empty Pandoc TikZ placeholders, found {len(placeholders)}")
    definitions = {
        f"Figure 4.1. {FIGURES[0]}": (
            (
                ("Provider", "STAC collections, CDSE metadata interface, or offline fixtures"),
                (
                    "Transport and validation",
                    "Injectable HTTP/OAuth2 transport, retries, pagination, and schema validation",
                ),
                (
                    "EOScene",
                    "Validated provider item transformed to a typed, transport-neutral record",
                ),
                (
                    "Provenance and catalog",
                    "SHA-256 raw payloads, append-only manifest, and SQLite SpatialCatalog",
                ),
                (
                    "Queries, API, and reports",
                    "Regional/quality filters, read-only educational API, and reproducible reports",
                ),
                (
                    "Scientific application",
                    "Local Planetary Computer COG raster analysis; CSC replay covers metadata only",
                ),
            ),
            "Provider metadata flows through validation, typed records, provenance, catalog, "
            "queries and reports to the scientific application. The raster case used local "
            "Planetary Computer assets; the CSC replay covered metadata only.",
        ),
        f"Figure 9.1. {WORKFLOW_FIGURE}": (
            WORKFLOW_STAGES,
            "Intended sequence from student onboarding through CSC access, catalogue discovery, "
            "dataset selection, retrieval, analysis, provenance, storage and reproduction. "
            "This is a teaching design, not a claim that every step ran on CSC.",
        ),
    }
    body = list(document.element.body)
    replacements = []
    for placeholder in placeholders:
        placeholder_index = body.index(placeholder._tbl)
        caption = next(
            (
                Paragraph(element, document)
                for element in body[placeholder_index + 1 :]
                if element.tag == qn("w:p") and Paragraph(element, document).text.strip()
            ),
            None,
        )
        definition_key = next(
            (key for key in definitions if caption is not None and caption.text.startswith(key)),
            None,
        )
        if caption is None or definition_key is None:
            raise ValueError(
                f"empty Pandoc TikZ placeholder is not followed by a recognized figure caption: "
                f"{'' if caption is None else caption.text}"
            )
        stages, description = definitions[definition_key]
        replacements.append((placeholder, caption, stages, description, definition_key))
    if {item[4] for item in replacements} != set(definitions):
        raise ValueError("each supported TikZ diagram must have exactly one placeholder")
    for placeholder, caption, stages, description, _ in replacements:
        diagram = _flow_diagram_table(document, stages, caption.text, description)
        caption._p.addprevious(diagram._tbl)
        placeholder._tbl.getparent().remove(placeholder._tbl)
    document.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pandoc", required=True, type=Path)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--csl", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=False)
    reference = args.output_dir.resolve() / "official_reference.docx"
    with ZipFile(args.template) as source, ZipFile(reference, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml"
                )
            target.writestr(item, data)
    output = args.output_dir.resolve() / "Thesis_Review_20261022.docx"
    subprocess.run(
        [
            str(args.pandoc.resolve()),
            "main.tex",
            "--from=latex",
            "--to=docx",
            "--citeproc",
            "--bibliography=references.bib",
            "--csl=" + str(args.csl.resolve()),
            "--reference-doc=" + str(reference),
            "-M",
            "title=",
            "-M",
            "author=",
            "-M",
            "date=",
            "--output=" + str(output),
        ],
        cwd=Path(__file__).parent,
        check=True,
    )
    counts = format_document(output)
    _insert_tikz_diagrams(output)
    if counts != {"figures": 7, "tables": 6}:
        raise ValueError(f"canonical caption inventory changed: {counts}; review export numbering")
    print(output)


if __name__ == "__main__":
    main()
